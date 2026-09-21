from __future__ import annotations

import os
import tempfile
import unittest
import zipfile
import json
import urllib.parse
from unittest.mock import patch
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from dashboard.delivery import DeliveryStore
from dashboard.cloud_state import VercelBlobState
from dashboard.metrics import DashboardStore, TARGET_KPIS
from dashboard.xlsx_reader import XlsxReader
from dashboard.xlsx_export import build_delivery_plan_xlsx
from dashboard.aging.exporter import build_aging_report_xlsx


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "MC_Dashboard_IMPORT.xlsx"
MASTER = ROOT / "data" / "delivery_master.json"


class DashboardCoreSmokeTests(unittest.TestCase):
    def test_bundled_workbook_contract(self):
        result = DashboardStore.validate(WORKBOOK)
        self.assertTrue(result.ok, result.message)

        store = DashboardStore(WORKBOOK)
        self.assertGreater(len(store.raw_records), 0)
        self.assertGreater(len(store.branches), 0)
        self.assertEqual(set(store.kpis), set(TARGET_KPIS))

        payload = store.bootstrap("guest")
        self.assertTrue(payload["has_data"])
        self.assertIn("status_summary", payload)
        self.assertIn("reorder_card", payload)

    def test_empty_kpi_series_are_not_shared(self):
        kpis = DashboardStore._empty_kpis()
        first, second = TARGET_KPIS[:2]
        kpis[first]["ytd"]["labels"].append("Jan")
        self.assertEqual(kpis[second]["ytd"]["labels"], [])
        self.assertEqual(kpis[first]["weekly"]["labels"], [])

    def test_xlsx_reader_reuses_sheet_parse(self):
        reader = XlsxReader(WORKBOOK)
        first = reader.read_sheet("Raw")
        second = reader.read_sheet("Raw")
        self.assertIs(first, second)
        self.assertGreater(len(first.rows), 2)


    def test_kpi_period_columns_are_dynamic_and_accept_text_dates(self):
        store = DashboardStore(None)
        target = "MUTI MC : Stock Outrate - Overall after PO Balance"
        rows = [
            [target, None, 0.10, 0.20, 0.30, 0.40],
            [None, None, 46053, 46081, "09/21/2026", "2026-10-31"],
        ]
        series = store._extract_kpi(rows, target)
        self.assertEqual(len(series["labels"]), 4)
        self.assertEqual(series["labels"][-2:], ["Sep 21, 2026", "Oct 31, 2026"])
        self.assertEqual(series["values"][-2:], [30.0, 40.0])

        # The supplied v2.46.9 baseline contains a newly-added YTD text-date
        # column. It must be visible instead of stopping at August.
        bundled = DashboardStore(WORKBOOK)
        for payload in bundled.kpis.values():
            self.assertEqual(payload["ytd"]["labels"][-1], "Sep 21, 2026")
            self.assertGreaterEqual(len(payload["ytd"]["labels"]), 9)

    def test_delivery_state_round_trip(self):
        store = DashboardStore(WORKBOOK)
        with tempfile.TemporaryDirectory() as temp_dir:
            delivery = DeliveryStore(MASTER, temp_dir)
            delivery.sync_dashboard_branches(store.raw_records)
            payload = delivery.bootstrap("Whole Week", store.raw_records)
            self.assertEqual(payload["days"], ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"])
            self.assertIn("summary", payload["analysis"])

            # Verify state writes remain readable after atomic replace.
            delivery.replace_allocations([
                {"model": "TEST MODEL", "branch": store.branches[0], "quantity": 1, "class": "A", "remarks": "smoke"}
            ])
            reloaded = DeliveryStore(MASTER, temp_dir)
            self.assertEqual(len(reloaded.allocations), 1)
            self.assertEqual(reloaded.allocations[0]["quantity"], 1.0)


    def test_delivery_redesign_keeps_required_control_ids(self):
        template = (ROOT / "templates" / "index.html").read_text(encoding="utf-8")
        required = [
            "deliveryTab", "deliveryDayFilter", "deliverySummaryCards", "deliveryWeekRibbon",
            "saveDeliveryPlanBtn", "exportDeliveryBtn", "clearDeliveryBoardBtn", "openDeliveryMasterBtn",
            "scheduleTripCount", "toggleScheduleBtn", "weeklyScheduleBody", "scheduleRows",
            "allocationCount", "toggleAllocationBtn", "allocationImportBody", "deliveryImportFile",
            "deliveryPlanPulse", "deliveryTruckGrid", "deliveryUnassigned", "deliveryPriorityList",
        ]
        for element_id in required:
            self.assertEqual(template.count(f'id="{element_id}"'), 1, element_id)
        self.assertIn("DELIVERY OPERATIONS CENTER", template)
        self.assertIn("Branch Dispatch Sequence", template)

    def test_model_aging_summary_has_professional_sort_controls(self):
        template = (ROOT / "templates" / "aging" / "dashboard.html").read_text(encoding="utf-8")
        aging_js = (ROOT / "static" / "aging" / "js" / "app.js").read_text(encoding="utf-8")
        self.assertIn('id="modelAgingSummaryTable"', template)
        self.assertEqual(template.count('class="table-sort-button'), 10)
        for index in range(1, 11):
            self.assertIn(f'data-sort-index="{index}"', template)
        self.assertIn('id="modelSortDesc"', template)
        self.assertIn('Highest → Lowest', template)
        self.assertIn('id="modelSortAsc"', template)
        self.assertIn('id="modelSortField"', template)
        self.assertIn('id="modelTableSearch"', template)
        self.assertIn('data-model-sort-preset="risk"', template)
        self.assertIn("sortRows('aged90','descending',{announce:false})", aging_js)
        self.assertIn("config[key].type==='number'?'descending':'ascending'", aging_js)

    def test_aging_model_risk_uses_same_exposure_thresholds(self):
        from dashboard.aging.analytics import _risk_level
        self.assertEqual(_risk_level(40)[0], "High")
        self.assertEqual(_risk_level(20)[0], "Watch")
        self.assertEqual(_risk_level(19.9)[0], "Controlled")


    def test_admin_access_session_contract(self):
        app_source = (ROOT / "app.py").read_text(encoding="utf-8")
        js = (ROOT / "static" / "js" / "dashboard-app.js").read_text(encoding="utf-8")
        aging_js = (ROOT / "static" / "js" / "aging-shell.js").read_text(encoding="utf-8")
        unified_js = (ROOT / "static" / "js" / "unified-import.js").read_text(encoding="utf-8")

        self.assertIn('SESSION_COOKIE_NAME"] = str(os.environ.get("SCM_SESSION_COOKIE_NAME") or "scm_idp_admin")', app_source)
        self.assertIn('@app.get("/api/session")', app_source)
        self.assertIn('"reauth_required": True', app_source)
        self.assertIn('session.permanent = True', app_source)
        self.assertIn("credentials:'same-origin'", js)
        self.assertIn("verifyAdminSession()", js)
        self.assertIn("/admin/export/management", js)
        self.assertIn("/admin/export/request", js)
        self.assertIn("/admin/delivery/plan", js)
        self.assertNotIn("window.location=`/admin/export/delivery", js)
        self.assertIn("credentials: 'same-origin'", aging_js)
        self.assertIn("xhr.withCredentials=true", unified_js)

    def test_delivery_export_builds_complete_week_workbook(self):
        store = DashboardStore(WORKBOOK)
        with tempfile.TemporaryDirectory() as temp_dir:
            delivery = DeliveryStore(MASTER, temp_dir)
            delivery.sync_dashboard_branches(store.raw_records)
            bundle = delivery.analysis_bundle(store.raw_records)
            workbook = build_delivery_plan_xlsx(
                "Whole Week", bundle["weekly"], bundle["daily"], delivery.schedule
            )

        with zipfile.ZipFile(BytesIO(workbook), "r") as archive:
            names = set(archive.namelist())
            for i in range(1, 8):
                self.assertIn(f"xl/worksheets/sheet{i}.xml", names)
            workbook_xml = archive.read("xl/workbook.xml").decode("utf-8")
            for day in ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday"]:
                self.assertIn(f'name="{day}"', workbook_xml)
            self.assertIn('name="Weekly Schedule"', workbook_xml)
            self.assertNotIn("<pane", archive.read("xl/worksheets/sheet1.xml").decode("utf-8"))

    def test_unified_import_tolerates_sheet_name_and_header_spacing(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            candidate = Path(temp_dir) / "renamed.xlsx"
            wb = load_workbook(WORKBOOK)
            wb["KPI_YTD_Input"].title = "KPI YTD Input"
            wb["KPI_WEEKLY_Input"].title = "KPI-WEEKLY-Input"
            wb["Raw"].insert_rows(1)
            wb["Raw"]["A1"] = "SCM CONTROL TOWER"
            wb["Aging"].insert_rows(1)
            wb["Aging"]["A1"] = "MOTORCYCLE AGING"
            wb.save(candidate)
            wb.close()

            result = DashboardStore.validate(candidate)
            self.assertTrue(result.ok, result.message)
            store = DashboardStore(candidate)
            self.assertGreater(len(store.raw_records), 0)
            self.assertGreater(len(store.management_records), 0)

    def test_unified_import_recovery_contract(self):
        app_source = (ROOT / "app.py").read_text(encoding="utf-8")
        import_js = (ROOT / "static" / "js" / "unified-import.js").read_text(encoding="utf-8")
        self.assertIn("candidate_store = DashboardStore(None)", app_source)
        self.assertIn("store.adopt_from(candidate_store, active)", app_source)
        self.assertIn("unified_import_errors.log", app_source)
        self.assertIn('"reference": reference', app_source)
        self.assertIn("parseXhrResponse", import_js)
        self.assertIn("Reference:", import_js)
        self.assertNotIn("xhr.responseType='json'", import_js)


    def test_delivery_cloud_persist_callback_contract(self):
        captured = []
        with tempfile.TemporaryDirectory() as temp_dir:
            def persist(path, payload):
                captured.append((Path(path).name, bytes(payload)))

            delivery = DeliveryStore(MASTER, temp_dir, persist_callback=persist)
            captured.clear()
            delivery.replace_allocations([
                {"model": "TEST MODEL", "branch": "TEST BRANCH", "quantity": 2, "class": "A"}
            ])
            self.assertTrue(captured)
            self.assertEqual(captured[-1][0], "delivery_allocations.json")
            self.assertIn(b"TEST MODEL", captured[-1][1])

    def test_cloud_core_revision_manifest_is_last_write(self):
        class MemoryBlobState(VercelBlobState):
            def __init__(self):
                super().__init__()
                self.objects = {}
                self.put_order = []

            @property
            def enabled(self):
                return True

            def _put_bytes(self, key, payload, *, content_type=None):
                full = self._key(key)
                self.objects[full] = bytes(payload)
                self.put_order.append(full)

            def get_bytes(self, key):
                return self.objects.get(self._key(key))

            def delete(self, key):
                self.objects.pop(self._key(key), None)

        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            workbook = root / "active_import.xlsx"
            aging = root / "aging.db"
            workbook.write_bytes(b"xlsx-state")
            aging.write_bytes(b"sqlite-state")

            cloud = MemoryBlobState()
            manifest = cloud.publish_core(workbook, aging, "IMP-TEST-001")

            self.assertEqual(cloud.put_order[-1], cloud._key(cloud.core_manifest_key))
            self.assertEqual(manifest["revision"], "IMP-TEST-001")

            runtime = root / "runtime"
            hydrated = cloud.hydrate_core(runtime)
            self.assertEqual(hydrated["revision"], "IMP-TEST-001")
            self.assertEqual((runtime / "active_import.xlsx").read_bytes(), b"xlsx-state")
            self.assertEqual((runtime / "aging" / "aging.db").read_bytes(), b"sqlite-state")

    def test_vercel_oidc_request_auth_is_supported_without_static_token(self):
        env = {
            "VERCEL": "1",
            "BLOB_STORE_ID": "store_abc123",
            "BLOB_READ_WRITE_TOKEN": "",
            "VERCEL_OIDC_TOKEN": "",
        }
        with patch.dict(os.environ, env, clear=False):
            cloud = VercelBlobState()
            self.assertFalse(cloud.enabled)
            cloud.bind_request_oidc("oidc-runtime-token")
            self.assertTrue(cloud.enabled)
            self.assertEqual(cloud.auth_mode, "oidc-request")
            headers = cloud._api_headers(content_type="application/json")
            self.assertEqual(headers["Authorization"], "Bearer oidc-runtime-token")
            self.assertEqual(headers["x-vercel-blob-store-id"], "abc123")
            self.assertEqual(headers["x-api-version"], "12")



    def test_vercel_session_readiness_contract(self):
        app_source = (ROOT / "app.py").read_text(encoding="utf-8")
        login = (ROOT / "templates" / "login.html").read_text(encoding="utf-8")
        dashboard_js = (ROOT / "static" / "js" / "dashboard-app.js").read_text(encoding="utf-8")
        self.assertIn('SESSION_KEY_SOURCE = "unknown"', app_source)
        self.assertIn('SERVERLESS_SESSION_READY', app_source)
        self.assertIn('"configuration_required": True', app_source)
        self.assertIn('"session_key_source": SESSION_KEY_SOURCE', app_source)
        self.assertIn('Deployment setup required.', login)
        self.assertIn('d.configuration_required', dashboard_js)

    def test_vercel_blob_http_probe_contract(self):
        objects = {}
        seen = []

        class Response:
            def __init__(self, payload=b"{}"):
                self.payload = payload
            def __enter__(self):
                return self
            def __exit__(self, *_args):
                return False
            def read(self):
                return self.payload

        def fake_urlopen(req, timeout=45):
            seen.append(req)
            method = req.get_method()
            parsed = urllib.parse.urlparse(req.full_url)
            if method == "PUT":
                pathname = urllib.parse.parse_qs(parsed.query)["pathname"][0]
                objects[pathname] = bytes(req.data or b"")
                return Response(b'{"ok":true}')
            if method == "GET":
                pathname = urllib.parse.unquote(parsed.path.lstrip("/"))
                return Response(objects[pathname])
            if method == "POST" and parsed.path.endswith("/delete"):
                payload = json.loads(bytes(req.data or b"{}").decode("utf-8"))
                for pathname in payload.get("urls", []):
                    objects.pop(pathname, None)
                return Response(b"{}")
            raise AssertionError(f"Unexpected request: {method} {req.full_url}")

        env = {
            "VERCEL": "1",
            "BLOB_STORE_ID": "store_abc123",
            "BLOB_READ_WRITE_TOKEN": "",
            "VERCEL_OIDC_TOKEN": "",
        }
        with patch.dict(os.environ, env, clear=False):
            cloud = VercelBlobState()
            cloud.bind_request_oidc("oidc-runtime-token")
            with patch("dashboard.cloud_state.urllib.request.urlopen", side_effect=fake_urlopen):
                result = cloud.probe()

        self.assertTrue(result["ok"])
        self.assertEqual(objects, {})
        self.assertEqual([req.get_method() for req in seen], ["PUT", "GET", "POST"])
        put_headers = {k.lower(): v for k, v in seen[0].header_items()}
        self.assertEqual(put_headers["authorization"], "Bearer oidc-runtime-token")
        self.assertEqual(put_headers["x-vercel-blob-store-id"], "abc123")
        self.assertEqual(put_headers["x-vercel-blob-access"], "private")
        self.assertEqual(put_headers["x-allow-overwrite"], "1")


    def test_vercel_runtime_fallback_does_not_require_blob(self):
        env = {
            "VERCEL": "1",
            "BLOB_STORE_ID": "",
            "BLOB_READ_WRITE_TOKEN": "",
            "VERCEL_OIDC_TOKEN": "",
            "SCM_REQUIRE_DURABLE_STORAGE": "0",
        }
        with patch.dict(os.environ, env, clear=False):
            cloud = VercelBlobState()
            cloud.bind_request_oidc("oidc-runtime-token")
            self.assertFalse(cloud.enabled)
            self.assertFalse(cloud.durable_required)
            self.assertTrue(cloud.allow_ephemeral)
            self.assertEqual(cloud.persistence_mode, "runtime-fallback")
            status = cloud.status()
            self.assertEqual(status.provider, "vercel-runtime-fallback")
            self.assertEqual(status.auth_mode, "runtime-fallback")

    def test_legacy_strict_storage_flags_do_not_block_universal_runtime(self):
        env = {
            "VERCEL": "1",
            "BLOB_STORE_ID": "",
            "BLOB_READ_WRITE_TOKEN": "",
            "VERCEL_OIDC_TOKEN": "",
            "SCM_REQUIRE_DURABLE_STORAGE": "1",
            "SCM_REQUIRE_BLOB_STORAGE": "1",
        }
        with patch.dict(os.environ, env, clear=False):
            cloud = VercelBlobState()
            cloud.bind_request_oidc("oidc-runtime-token")
            self.assertFalse(cloud.enabled)
            self.assertFalse(cloud.durable_required)
            self.assertTrue(cloud.allow_ephemeral)
            self.assertEqual(cloud.status().provider, "vercel-runtime-fallback")

    def test_vercel_persistent_storage_contract(self):
        app_source = (ROOT / "app.py").read_text(encoding="utf-8")
        cloud_source = (ROOT / "dashboard" / "cloud_state.py").read_text(encoding="utf-8")
        requirements = (ROOT / "requirements.txt").read_text(encoding="utf-8")

        # Vercel must always use the writable system temp area, even when an old
        # SCM_DATA_DIR environment variable survives from a previous deployment.
        self.assertIn('if IS_VERCEL:\n    os.environ["SCM_DATA_DIR"] = str(Path(tempfile.gettempdir()) / "scm-idp-dashboard")', app_source)
        self.assertIn('IMPORT_TMP_DIR = Path(tempfile.gettempdir()) / "scm-idp-dashboard-imports"', app_source)
        self.assertIn('dir=IMPORT_TMP_DIR', app_source)
        self.assertNotIn('dir=UPLOAD_DIR)\n    os.close(fd)\n    path = Path(raw_path)\n    try:\n        file_storage.save(path)', app_source)

        self.assertIn('request.headers.get("x-vercel-oidc-token", "")', app_source)
        self.assertIn('cloud_state.probe()', app_source)
        self.assertIn('cloud_state.publish_core(active, aging_cloud_snapshot, reference)', app_source)
        self.assertNotIn('error_stage="checking Vercel storage"', app_source)
        self.assertIn('runtime-fallback', app_source)
        self.assertIn('strict_durable_storage', app_source)
        self.assertIn('cloud_state.publish_empty_core(reset_reference)', app_source)
        self.assertIn('persist_callback=_persist_small_state if cloud_state.enabled else None', app_source)
        self.assertIn('delivery_store.set_persist_callback(_persist_small_state)', app_source)

        self.assertIn('x-vercel-oidc-token', cloud_source)
        self.assertIn('BLOB_STORE_ID', cloud_source)
        self.assertIn('x-vercel-blob-store-id', cloud_source)
        self.assertIn('x-vercel-blob-access', cloud_source)
        self.assertIn('core/current.json', cloud_source)
        # v2.46.7 keeps the Blob HTTP API directly, removing SDK-version coupling.
        self.assertNotIn('vercel>=', requirements.lower())


    def test_aging_unit_traceability_has_filters_and_full_dataset_export(self):
        template = (ROOT / "templates" / "aging" / "dashboard.html").read_text(encoding="utf-8")
        routes = (ROOT / "dashboard" / "aging" / "routes.py").read_text(encoding="utf-8")
        aging_js = (ROOT / "static" / "aging" / "js" / "app.js").read_text(encoding="utf-8")
        for element_id in ["unitTrace", "unitTraceForm", "unitTraceSearch", "unitAgeFilter", "unitSortFilter"]:
            self.assertEqual(template.count(f'id="{element_id}"'), 1, element_id)
        self.assertIn('data-unit-band="critical"', template)
        self.assertIn('Highest Value → Lowest', template)
        self.assertIn('@aging_bp.get("/export.xlsx")', routes)
        self.assertIn('apply_unit_filters(rows, unit_filters)', routes)
        self.assertIn("unitForm.requestSubmit()", aging_js)

    def test_aging_export_builds_professional_multisheet_workbook(self):
        summary = {
            "total_qty": 10, "total_value": 500000, "avg_age": 150, "aged_90": 5, "aged_90_pct": 50.0,
            "aged_90_value": 250000, "aged_180": 3, "aged_365": 1, "healthy_qty": 5, "oldest": 430,
            "branch_count": 2, "area_count": 1, "risk_level": "High", "risk_message": "Test exposure",
            "detailed_bucket_labels": ["0-30 DAYS", "31-60 DAYS", "61-90 DAYS", "91-180 DAYS", "181-365 DAYS", "366+ DAYS"],
            "detailed_bucket_values": [1, 2, 2, 2, 2, 1],
            "top_area": {"name": "AREA I", "aged_90": 5, "aged_90_pct": 50, "aged_value": 250000, "oldest": 430, "branch_count": 2},
            "area_ranking_all": [{"name": "AREA I", "risk_level": "High", "branch_count": 2, "qty": 10, "value": 500000, "avg_age": 150, "aged_90": 5, "aged_90_pct": 50, "aged_value": 250000, "aged_180": 3, "aged_365": 1, "oldest": 430}],
            "top_branch": {"name": "BRANCH A", "area": "AREA I", "aged_90": 3, "aged_90_pct": 60, "aged_value": 150000, "oldest": 430},
            "top_model": {"standard_description": "MODEL X", "aged_90": 4, "aged_90_pct": 80, "aged_value": 200000, "oldest": 430, "branch_count": 2, "area_count": 1},
            "model_summary_all": [{"standard_description": "MODEL X", "risk_level": "High", "qty": 5, "avg_age": 220, "aged_90": 4, "aged_90_pct": 80, "aged_value": 200000, "oldest": 430, "branch_count": 2, "area_count": 1}],
        }
        units = [{"branch_name": "BRANCH A", "area": "AREA I", "standard_description": "MODEL X", "brand": "HONDA", "engine_no": "E1", "chassis": "C1", "created_on": "2025-07-18", "incoming_date": "2025-07-10", "age_days": 430, "qty": 1, "amount": 50000, "inventory_value": 50000, "location": "SHOWROOM", "barcode": "BC1"}]
        payload = build_aging_report_xlsx(summary=summary, unit_rows=units, as_of="2026-09-21", basis="branch", filters={}, unit_filters={"age": "all", "sort": "oldest", "q": ""}, source_filename="MC.xlsx")
        with zipfile.ZipFile(BytesIO(payload), "r") as archive:
            self.assertIsNone(archive.testzip())
            self.assertIn("xl/worksheets/sheet5.xml", archive.namelist())
        with tempfile.NamedTemporaryFile(suffix=".xlsx") as temp:
            temp.write(payload); temp.flush()
            wb = load_workbook(temp.name, read_only=False)
            self.assertEqual(wb.sheetnames, ["Executive Summary", "Area Intelligence", "Model Intelligence", "Unit Detail", "Data Dictionary"])
            self.assertEqual(wb["Area Intelligence"].freeze_panes, "C6")
            self.assertEqual(wb["Area Intelligence"].auto_filter.ref, "A5:M6")
            self.assertEqual(wb["Area Intelligence"]["B6"].value, "AREA I")
            self.assertEqual(wb["Model Intelligence"].freeze_panes, "C5")
            self.assertEqual(wb["Unit Detail"].freeze_panes, "C6")
            self.assertEqual(wb["Unit Detail"].auto_filter.ref, "A5:N6")
            wb.close()


    def test_ytd_chart_keeps_months_compact_and_marks_latest_date(self):
        js = (ROOT / "static" / "js" / "dashboard-app.js").read_text(encoding="utf-8")
        self.assertIn("i===labels.length-1?shortCurrentPeriodLabel(label):monthOnly(label)", js)
        self.assertIn("Latest data", js)
        self.assertIn("longPeriodLabel(d.labels[d.labels.length-1])", js)
        bundled = DashboardStore(WORKBOOK)
        latest = next(iter(bundled.kpis.values()))["ytd"]["labels"][-1]
        self.assertEqual(latest, "Sep 21, 2026")

    def test_v2474_theme_readability_guard_is_loaded_everywhere(self):
        guard = (ROOT / "static" / "css" / "v2474-theme-readability.css").read_text(encoding="utf-8")
        for rel in ["templates/index.html", "templates/aging/base.html", "templates/login.html"]:
            self.assertIn("v2474-theme-readability.css", (ROOT / rel).read_text(encoding="utf-8"), rel)
        self.assertIn('html[data-theme="light"] [class~="text-slate-500"]', guard)
        self.assertIn('html[data-theme="light"] .bg-blue-600', guard)
        self.assertIn('.kpi-latest-date', guard)

    def test_v2475_theme_visibility_guard_and_custom_trend_legend(self):
        guard = (ROOT / "static" / "css" / "v2475-visibility-fix.css").read_text(encoding="utf-8")
        js = (ROOT / "static" / "js" / "dashboard-app.js").read_text(encoding="utf-8")
        for rel in ["templates/index.html", "templates/aging/base.html", "templates/login.html"]:
            self.assertIn("v2475-visibility-fix.css", (ROOT / rel).read_text(encoding="utf-8"), rel)
        self.assertIn('body.aging-module .kpi-value{color:#0b1f3a!important}', guard)
        self.assertIn('.kpi-legend-line.trend', guard)
        self.assertIn("const trendIcon=trendDirection==='up'?'↗':trendDirection==='down'?'↘':'→'", js)
        self.assertIn('data-kpi-chart-legend', js)
        self.assertIn("legend:{display:false}", js)
        self.assertIn("trend:'#0369a1'", js)

    def test_aging_unit_filter_preserves_viewport_across_reload(self):
        aging_js = (ROOT / "static" / "aging" / "js" / "app.js").read_text(encoding="utf-8")
        self.assertIn("scm-aging-unit-view-v1", aging_js)
        self.assertIn("rememberUnitViewport()", aging_js)
        self.assertIn("restoreUnitViewport", aging_js)
        self.assertIn("viewportTop:unitForm.getBoundingClientRect().top", aging_js)
        self.assertIn("window.scrollBy(0,delta)", aging_js)
        self.assertIn("focus({preventScroll:true})", aging_js)

    def test_theme_contrast_guard_is_loaded_for_main_aging_and_login(self):
        guard = (ROOT / "static" / "css" / "v2471-theme-contrast.css").read_text(encoding="utf-8")
        for rel in ["templates/index.html", "templates/aging/base.html", "templates/login.html"]:
            self.assertIn("v2471-theme-contrast.css", (ROOT / rel).read_text(encoding="utf-8"), rel)
        self.assertIn('html[data-theme="dark"] [class*="text-slate-600"]', guard)
        self.assertIn('html[data-theme="light"] [class*="text-amber-300"]', guard)
        self.assertIn('.sidebar-logo-img,.sidebar-logo-mini', guard)
        self.assertIn('filter:none!important', guard)

if __name__ == "__main__":
    unittest.main()
