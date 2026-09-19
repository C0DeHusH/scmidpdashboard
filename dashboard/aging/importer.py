from __future__ import annotations

import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import transaction

HEADER_ALIASES = {
    "BRANCH": ["BRANCH"],
    "AREA": ["AREA", "REGION", "DISTRICT"],
    "INCOMING DATE": ["INCOMING DATE", "RECEIVED DATE", "DATE RECEIVED"],
    "CREATED ON": ["CREATED ON", "CREATED DATE", "BRANCH DATE"],
    "BARCODE": ["BARCODE", "ITEM CODE", "MODEL CODE"],
    "DESCRIPTION": ["DESCRIPTION", "ITEM DESCRIPTION"],
    "QTY": ["QTY", "QUANTITY"],
    "STANDARD DESCRIPTION": ["STANDARD DESCRIPTION", "STANDARD DESC", "MODEL"],
    "AMOUNT": ["AMOUNT", "UNIT COST", "COST"],
    "COMPANY": ["COMPANY"],
    "ENGINE NO.": ["ENGINE NO.", "ENGINE NO", "ENGINE NUMBER"],
    "CHASS.": ["CHASS.", "CHASSIS", "CHASSIS NO.", "CHASSIS NO"],
    "LOCATION": ["LOCATION", "STOCK LOCATION"],
    "BRAND": ["BRAND"],
    "COLOR": ["COLOR", "COLOUR"],
    # Workbook aging formula columns are intentionally not authoritative.
    # Python recalculates aging from the selected/auto-detected report date.
    "AGING DAYS BRANCH": ["AGING DAYS BRANCH"],
    "AGING DAYS COMPANY": ["AGING DAYS COMPANY"],
    "AGING BRANCH": ["AGING BRANCH"],
    "AGING COMPANY": ["AGING COMPANY"],
}


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).upper()


def _canonical_branch(value: Any) -> str:
    """Normalize branch aliases shared by the consolidated SCM + Aging workbook.

    Aging historically uses M1 / M2 prefixes while Raw uses MUTI. Honda branches
    can also be prefixed by M1. Removing only the organization prefix lets both
    sources resolve to one branch/area master without hard-coded branch lists.
    """
    text = _norm(value)
    text = re.sub(r"^(?:M1|M2|MUTI)\s+", "", text).strip()
    return text


def _is_good(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, float) and math.isnan(value):
        return False
    text = str(value).strip()
    return bool(text) and not text.startswith("#") and not text.startswith("=")


def _clean_text(value: Any) -> str:
    if not _is_good(value):
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def _as_number(value: Any, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        text = re.sub(r"[^0-9.\-]", "", str(value))
        try:
            return float(text) if text else default
        except ValueError:
            return default


def _as_date(value: Any) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%m/%d/%y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def _location_branch(location: Any) -> str:
    text = _clean_text(location).upper()
    if not text:
        return "UNMAPPED"
    token = re.split(r"[/\\\-\s]", text, maxsplit=1)[0].strip()
    return token or "UNMAPPED"


def _header_map(header_values: list[Any]) -> dict[str, int]:
    normalized = {_norm(v): i for i, v in enumerate(header_values) if _norm(v)}
    result: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            if _norm(alias) in normalized:
                result[canonical] = normalized[_norm(alias)]
                break
    return result


def _value(row: tuple[Any, ...], mapping: dict[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _find_header_row(ws, required: set[str], max_rows: int = 8) -> tuple[int, list[Any]] | None:
    for row_no, row in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        normalized = {_norm(v) for v in row if _norm(v)}
        if required.issubset(normalized):
            return row_no, list(row)
    return None


def _choose_aging_sheet(workbook, requested: str | None = None):
    if requested and requested in workbook.sheetnames:
        return workbook[requested]
    for preferred in ("Aging", "AGING", "Motorcycle Aging"):
        if preferred in workbook.sheetnames:
            return workbook[preferred]
    required = {"BRANCH", "INCOMING DATE", "CREATED ON", "STANDARD DESCRIPTION"}
    for name in workbook.sheetnames:
        ws = workbook[name]
        if _find_header_row(ws, required):
            return ws
    raise ValueError("Unified import is missing an Aging worksheet with Branch, Incoming Date, Created On and Standard Description columns.")


def _raw_branch_area_master(workbook) -> dict[str, tuple[str, str]]:
    """Build a consolidated branch -> area lookup from the Raw worksheet."""
    if "Raw" not in workbook.sheetnames:
        return {}
    ws = workbook["Raw"]
    found = _find_header_row(ws, {"BRANCH", "AREA"}, max_rows=8)
    if not found:
        return {}
    header_row, headers = found
    first_index: dict[str, int] = {}
    for i, value in enumerate(headers):
        key = _norm(value)
        if key and key not in first_index:  # Raw can contain another Area later.
            first_index[key] = i
    branch_idx = first_index.get("BRANCH")
    area_idx = first_index.get("AREA")
    if branch_idx is None or area_idx is None:
        return {}
    lookup: dict[str, tuple[str, str]] = {}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        branch = _clean_text(row[branch_idx] if branch_idx < len(row) else None)
        area = _clean_text(row[area_idx] if area_idx < len(row) else None)
        if not branch:
            continue
        key = _canonical_branch(branch)
        current = lookup.get(key)
        if not current or (not current[1] and area):
            lookup[key] = (branch, area)
    return lookup


def validate_unified_aging(path: str | Path, sheet_name: str | None = "Aging") -> dict[str, Any]:
    """Fast structural validation used before the unified system import commits."""
    wb = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        ws = _choose_aging_sheet(wb, sheet_name)
        found = _find_header_row(ws, {"BRANCH", "INCOMING DATE", "CREATED ON", "STANDARD DESCRIPTION"})
        if not found:
            raise ValueError("Aging sheet is missing required columns: Branch, Incoming Date, Created On, Standard Description.")
        header_row, headers = found
        mapping = _header_map(headers)
        return {
            "sheet": ws.title,
            "header_row": header_row,
            "has_area": "AREA" in mapping,
            "raw_area_lookup": bool(_raw_branch_area_master(wb)),
        }
    finally:
        wb.close()


def import_excel(
    path: str | Path,
    filename: str,
    as_of_date: str | None = None,
    mode: str = "replace",
    sheet_name: str | None = "Aging",
) -> dict[str, Any]:
    """Import Aging data from the same consolidated workbook used by SCM.

    v2.44 unified behavior:
    * Prefer the worksheet named ``Aging`` instead of assuming the first sheet.
    * AREA is optional on Aging. When absent, it is resolved from Raw Branch/Area.
    * M1/M2/MUTI branch prefixes are reconciled automatically.
    * Branch names are standardized to the Raw-sheet branch master when matched.
    * Excel aging formula columns are ignored; Python recalculates from source dates.
    * The latest Created On / Incoming Date is the report date when none is supplied.
    """
    path = Path(path)
    wb_values = load_workbook(path, read_only=True, data_only=True)
    wb_formula = load_workbook(path, read_only=True, data_only=False)
    try:
        ws_v = _choose_aging_sheet(wb_values, sheet_name)
        ws_f = wb_formula[ws_v.title]
        found = _find_header_row(ws_v, {"BRANCH", "INCOMING DATE", "CREATED ON", "STANDARD DESCRIPTION"})
        if not found:
            raise ValueError("Aging sheet is missing required columns: Branch, Incoming Date, Created On, Standard Description.")
        header_row, header_v = found
        mapping = _header_map(header_v)
        required = ["BRANCH", "INCOMING DATE", "CREATED ON", "STANDARD DESCRIPTION"]
        missing = [x for x in required if x not in mapping]
        if missing:
            raise ValueError("Aging sheet is missing required column(s): " + ", ".join(missing))

        raw_master = _raw_branch_area_master(wb_values)
        if "AREA" not in mapping and not raw_master:
            raise ValueError("Aging does not contain Area and the Raw Branch/Area master could not be read.")

        explicit_as_of: date | None = None
        if as_of_date:
            try:
                explicit_as_of = datetime.strptime(as_of_date, "%Y-%m-%d").date()
            except ValueError as exc:
                raise ValueError("As-of date must use YYYY-MM-DD format.") from exc

        values_iter = ws_v.iter_rows(min_row=header_row + 1, values_only=True)
        formula_iter = ws_f.iter_rows(min_row=header_row + 1, values_only=True)
        staged_rows: list[dict[str, Any]] = []
        latest_source_date: date | None = None
        branch_seed: dict[str, tuple[str, str]] = {}
        skipped = 0
        fallback_branch = 0
        fallback_description = 0
        fallback_area = 0
        standardized_branches = 0

        for source_row, (row_v, row_f) in enumerate(zip(values_iter, formula_iter), start=header_row + 1):
            description = _clean_text(_value(row_v, mapping, "DESCRIPTION")) or _clean_text(_value(row_f, mapping, "DESCRIPTION"))
            location = _clean_text(_value(row_v, mapping, "LOCATION")) or _clean_text(_value(row_f, mapping, "LOCATION"))
            engine = _clean_text(_value(row_v, mapping, "ENGINE NO.")) or _clean_text(_value(row_f, mapping, "ENGINE NO."))
            chassis = _clean_text(_value(row_v, mapping, "CHASS.")) or _clean_text(_value(row_f, mapping, "CHASS."))
            barcode = _clean_text(_value(row_v, mapping, "BARCODE")) or _clean_text(_value(row_f, mapping, "BARCODE"))
            if not any((description, location, engine, chassis, barcode)):
                skipped += 1
                continue

            branch_raw = _clean_text(_value(row_v, mapping, "BRANCH")) or _clean_text(_value(row_f, mapping, "BRANCH"))
            if not branch_raw:
                branch_raw = _location_branch(location)
                fallback_branch += 1

            master_match = raw_master.get(_canonical_branch(branch_raw))
            if master_match:
                master_branch, master_area = master_match
                if master_branch and _norm(master_branch) != _norm(branch_raw):
                    standardized_branches += 1
                branch_display = master_branch or branch_raw
            else:
                master_area = ""
                branch_display = branch_raw

            branch_key = _norm(branch_display) or "UNMAPPED"
            area_raw = _clean_text(_value(row_v, mapping, "AREA")) or _clean_text(_value(row_f, mapping, "AREA"))
            if not area_raw:
                area_raw = master_area or "UNMAPPED"
                if area_raw == "UNMAPPED":
                    fallback_area += 1

            std_desc = _clean_text(_value(row_v, mapping, "STANDARD DESCRIPTION")) or _clean_text(_value(row_f, mapping, "STANDARD DESCRIPTION"))
            if not std_desc:
                std_desc = description
                fallback_description += 1

            incoming_dt = _as_date(_value(row_v, mapping, "INCOMING DATE")) or _as_date(_value(row_f, mapping, "INCOMING DATE"))
            created_dt = _as_date(_value(row_v, mapping, "CREATED ON")) or _as_date(_value(row_f, mapping, "CREATED ON"))
            for candidate in (incoming_dt, created_dt):
                if candidate and (latest_source_date is None or candidate > latest_source_date):
                    latest_source_date = candidate

            branch_seed[branch_key] = (branch_display or branch_key, area_raw)
            staged_rows.append({
                "source_row": source_row,
                "branch_key": branch_key,
                "branch_original": branch_display or branch_key,
                "area": area_raw,
                "incoming_date": incoming_dt.isoformat() if incoming_dt else None,
                "created_on": created_dt.isoformat() if created_dt else None,
                "barcode": barcode,
                "description": description,
                "qty": _as_number(_value(row_v, mapping, "QTY"), 1.0),
                "standard_description": std_desc,
                "amount": _as_number(_value(row_v, mapping, "AMOUNT"), 0.0),
                "company": _clean_text(_value(row_v, mapping, "COMPANY")) or _clean_text(_value(row_f, mapping, "COMPANY")),
                "engine_no": engine,
                "chassis": chassis,
                "location": location,
                "brand": _clean_text(_value(row_v, mapping, "BRAND")) or _clean_text(_value(row_f, mapping, "BRAND")),
                "color": _clean_text(_value(row_v, mapping, "COLOR")) or _clean_text(_value(row_f, mapping, "COLOR")),
            })

        if not staged_rows:
            raise ValueError("No motorcycle inventory rows were found in the Aging worksheet.")

        final_as_of = explicit_as_of or latest_source_date or date.today()
        if latest_source_date and final_as_of < latest_source_date:
            raise ValueError(
                f"As-of date {final_as_of.isoformat()} is earlier than the newest inventory date "
                f"{latest_source_date.isoformat()}. Use Auto-detect or choose a later date."
            )

        rows_to_insert = [(
            r["source_row"], r["branch_key"], r["branch_original"], r["area"],
            r["incoming_date"], r["created_on"], r["barcode"], r["description"], r["qty"],
            r["standard_description"], r["amount"], r["company"], r["engine_no"], r["chassis"],
            r["location"], r["brand"], r["color"], final_as_of.isoformat(),
        ) for r in staged_rows]

        with transaction() as conn:
            if mode == "replace":
                conn.execute("DELETE FROM units")
                conn.execute("DELETE FROM imports")
                conn.execute("DELETE FROM branch_area")

            cur = conn.execute(
                "INSERT INTO imports(filename, as_of_date, row_count, import_mode) VALUES(?,?,?,?)",
                (filename, final_as_of.isoformat(), len(rows_to_insert), mode),
            )
            import_id = cur.lastrowid

            for key, (name, area) in branch_seed.items():
                conn.execute(
                    """
                    INSERT INTO branch_area(branch_key, branch_name, area, updated_at)
                    VALUES(?,?,?,CURRENT_TIMESTAMP)
                    ON CONFLICT(branch_key) DO UPDATE SET
                        branch_name=excluded.branch_name,
                        area=CASE WHEN excluded.area<>'UNMAPPED' THEN excluded.area ELSE branch_area.area END,
                        updated_at=CURRENT_TIMESTAMP
                    """,
                    (key, name or key, area or "UNMAPPED"),
                )

            conn.executemany(
                """
                INSERT INTO units(
                    import_id, source_row, branch_key, branch_original, area, incoming_date, created_on,
                    barcode, description, qty, standard_description, amount, company, engine_no,
                    chassis, location, brand, color, imported_as_of
                ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                """,
                [(import_id, *r) for r in rows_to_insert],
            )

        return {
            "rows": len(rows_to_insert),
            "skipped": skipped,
            "branches": len(branch_seed),
            "areas": len({r["area"] for r in staged_rows if r["area"] and r["area"] != "UNMAPPED"}),
            "branch_fallbacks": fallback_branch,
            "area_fallbacks": fallback_area,
            "description_fallbacks": fallback_description,
            "standardized_branches": standardized_branches,
            "sheet": ws_v.title,
            "as_of_date": final_as_of.isoformat(),
            "as_of_source": "manual" if explicit_as_of else "auto",
            "latest_source_date": latest_source_date.isoformat() if latest_source_date else None,
        }
    finally:
        wb_values.close()
        wb_formula.close()
