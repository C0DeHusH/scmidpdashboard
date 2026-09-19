from __future__ import annotations

from io import BytesIO
from typing import Any, Dict, List
import math
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from functools import lru_cache

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from PIL import Image, ImageDraw, ImageFont

DARK = RGBColor(4, 14, 31)
DARK2 = RGBColor(8, 28, 52)
SKY_BLUE = RGBColor(96, 165, 250)  # Primary highlight
WHITE = RGBColor(248, 251, 255)
MUTED = RGBColor(157, 181, 207)
RED = RGBColor(248, 113, 113)
GREEN = RGBColor(74, 222, 128)
TEAL = RGBColor(56, 189, 248)
LOGO_PATH = Path(__file__).resolve().parents[1] / "static" / "brilliant4_logo.png"
NAVY = RGBColor(3, 12, 28)
NAVY_PANEL = RGBColor(7, 30, 57)
BLUE = RGBColor(37, 99, 235)
BLUE_LINE = RGBColor(59, 130, 246)
CYAN = RGBColor(56, 189, 248)
AQUA = RGBColor(125, 211, 252)
BLUE_PALE = RGBColor(191, 219, 254)



def _whole(value):
    if value is None or value == "—":
        return value
    try:
        return int(Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    except Exception:
        return value


def _doi_whole(value):
    if value is None or value == "—":
        return value
    try:
        return int(math.ceil(max(0.0, float(value))))
    except Exception:
        return value


def _delta_favorable(delta, good):
    if delta is None:
        return None
    try:
        d = float(delta)
    except Exception:
        return None
    if abs(d) < 1e-12 or good == "balanced":
        return None
    return d > 0 if good == "high" else d < 0


def _set_bg(slide, color=DARK):
    fill = slide.background.fill
    fill.solid(); fill.fore_color.rgb = color

def _generated_date(data: Dict[str, Any]) -> str:
    """Cover-page date: when this export was generated, not the latest KPI date."""
    raw = (data or {}).get("export_generated_at")
    try:
        dt = datetime.fromisoformat(str(raw)) if raw else datetime.now()
    except Exception:
        dt = datetime.now()
    return dt.strftime("%B %d, %Y").replace(" 0", " ")


def _kpi_heading(label: str):
    u = (label or "").upper()
    if "BEFORE PO" in u:
        return "STOCK-OUT", "• BEFORE PO BALANCE", "VISIBILITY  •  CONTROL  •  SUPPLY RISK  •  A STRONGER TOMORROW"
    if "AFTER PO" in u:
        return "STOCK-OUT", "• AFTER PO BALANCE", "FULFILMENT  •  CONTROL  •  SUPPLY RECOVERY  •  A STRONGER TOMORROW"
    if "PER BRANCH" in u:
        return "STOCK-OUT", "• PER BRANCH", "BRANCH VISIBILITY  •  PRIORITY  •  SERVICE LEVEL  •  GROWTH"
    if "CLASS A" in u and "DOI" in u:
        return "CLASS A DOI", "• AVAILABILITY", "PRIORITY MODELS  •  COVERAGE  •  REPLENISHMENT  •  GROWTH"
    if "CLASS A" in u and "STOCK" in u:
        return "CLASS A STOCK-OUT", "• NETWORK PRIORITY", "CLASS A VISIBILITY  •  CONTROL  •  REPLENISHMENT  •  GROWTH"
    if "DOI" in u:
        return "DAYS OF INVENTORY", "• NETWORK COVERAGE", "AVAILABILITY  •  BALANCE  •  INVENTORY HEALTH  •  GROWTH"
    return label.upper(), "", "VISIBILITY  •  CONTROL  •  SUPPLY  •  GROWTH"


def _kpi_actions(label: str):
    u = (label or "").upper()
    if "STOCK" in u:
        return [
            ("01", "REDUCE STOCK-OUTS", "Protect demand before PO arrival"),
            ("02", "STRENGTHEN PLANNING", "Escalate supply and back-order actions"),
            ("03", "DRIVE CONVERSION", "Use Class B & C alternatives when needed"),
            ("04", "ENABLE GROWTH", "Support sales with supply visibility"),
        ]
    return [
        ("01", "PROTECT AVAILABILITY", "Sustain healthy days of inventory"),
        ("02", "REBALANCE STOCK", "Move units to priority branches and models"),
        ("03", "FOCUS CLASS A", "Prioritize high-rank, low-DoI models"),
        ("04", "ENABLE GROWTH", "Support sales with inventory visibility"),
    ]


def _kpi_chart_png(labels: List[str], values: List[float], trend: List[float], percent=False, period="ytd", good="low") -> BytesIO:
    """Return a cached in-memory KPI chart image.

    The previous implementation rendered to undeleted temporary PNG files for
    every export. In-memory rendering removes disk I/O and the LRU cache reuses
    charts while the imported data is unchanged.
    """
    payload = _render_kpi_chart_png(
        tuple(str(x) for x in labels),
        tuple(values),
        tuple(trend),
        bool(percent),
        str(period),
        str(good),
    )
    stream = BytesIO(payload)
    stream.seek(0)
    return stream


@lru_cache(maxsize=64)
def _render_kpi_chart_png(labels, values, trend, percent=False, period="ytd", good="low") -> bytes:
    fig, ax = plt.subplots(figsize=(10.4, 3.25), dpi=150)
    fig.patch.set_alpha(0)
    ax.set_facecolor((0, 0, 0, 0))
    for spine in ax.spines.values():
        spine.set_color("#527aa5")
        spine.set_alpha(0.55)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)

    vals = [_whole(v) for v in values]
    tr = [float(v) for v in trend]
    x = list(range(len(vals)))
    ax.plot(x, vals, marker="o", markersize=6.5, linewidth=3.2, color="#3b82f6", label="Actual", zorder=4)
    if vals:
        ax.fill_between(x, vals, color="#3b82f6", alpha=0.15, zorder=1)

    numeric = [v for v in vals + tr if isinstance(v, (int, float))]
    spread = max(numeric) - min(numeric) if numeric else 1
    gap = max(1, spread * 0.10)
    required = max([((a + gap) - t) for a, t in zip(vals, tr) if isinstance(a, (int, float)) and isinstance(t, (int, float))] or [0])
    display_trend = [t + max(0, required) for t in tr]
    ax.plot(x, display_trend, linestyle=(0, (6, 4)), linewidth=2.5, color="#93c5fd", label="Trend Direction", zorder=3)

    data_gap = max(0.7, spread * 0.045)
    for xi, yi in zip(x, vals):
        if isinstance(yi, (int, float)):
            txt = f"{yi}%" if percent else f"{yi}"
            ax.text(xi, yi + data_gap, txt, color="#bfdbfe", fontsize=8.8, fontweight="bold", ha="center", va="bottom", zorder=5)

    if labels:
        shown = []
        for raw in labels:
            try:
                dt = datetime.strptime(str(raw), "%b %d, %Y")
                if period == "ytd":
                    shown.append(dt.strftime("%b") if dt.day not in (14, 15) else dt.strftime("%b %d").replace(" 0", " "))
                else:
                    shown.append(dt.strftime("%b %d").replace(" 0", " "))
            except Exception:
                shown.append(str(raw))
        idxs = list(range(len(labels)))
        ax.set_xticks(idxs)
        ax.set_xticklabels([shown[i] for i in idxs], fontsize=8.5, color="#dbeafe")

    allv = [v for v in vals + display_trend if isinstance(v, (int, float))]
    if allv:
        lo, hi = min(allv), max(allv)
        if percent:
            lower = min(0, math.floor((lo - max(5, spread * .14)) / 10) * 10)
            upper = max(10, math.ceil((hi + max(5, spread * .14)) / 10) * 10)
        else:
            lower = max(0, math.floor(lo - max(3, spread * .12)))
            upper = math.ceil(hi + max(3, spread * .15))
        if upper <= lower:
            upper = lower + 10
        ax.set_ylim(lower, upper)

    ax.grid(True, axis="both", linestyle="--", linewidth=0.65, alpha=0.18, color="#4f78a2")
    ax.tick_params(colors="#b8d0e8", labelsize=8.5)
    from matplotlib.ticker import FuncFormatter
    if percent:
        ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{int(y)}%"))
    else:
        ax.yaxis.set_major_formatter(FuncFormatter(lambda y, _: f"{int(y)}"))
    ax.legend(loc="upper right", bbox_to_anchor=(1.0, 1.12), frameon=False, labelcolor="#dbeafe", fontsize=8.5, ncol=2, handlelength=3)
    fig.subplots_adjust(left=0.07, right=0.99, top=0.88, bottom=0.16)
    out = BytesIO()
    fig.savefig(out, format="png", transparent=True)
    plt.close(fig)
    return out.getvalue()


def _add_kpi_template_slide(prs, blank, name: str, kpi: Dict[str, Any], period_key: str, period_title: str):
    slide=prs.slides.add_slide(blank)
    _set_bg(slide, NAVY)
    # Native vector background: lighter, faster and fully editable.
    top_rule=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.075))
    top_rule.fill.solid(); top_rule.fill.fore_color.rgb=BLUE; top_rule.line.fill.background()
    glow=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.55), Inches(0.075), Inches(2.78), Inches(0.18))
    glow.fill.solid(); glow.fill.fore_color.rgb=CYAN; glow.line.fill.background()

    meta=kpi.get("meta",{})
    label=meta.get("label",name)
    unit=meta.get("unit")
    good=meta.get("good","low")
    d=kpi.get(period_key,{}) or {}
    values=list(d.get("values",[]) or [])
    labels=list(d.get("labels",[]) or [])
    trend=list(d.get("trend",[]) or [])
    latest=d.get("latest")
    delta=d.get("delta")
    primary,secondary,strap=_kpi_heading(label)

    # Header
    secondary_x = min(6.05, 0.48 + 0.22 * len(primary))
    _textbox(slide,0.46,0.66,max(2.3,secondary_x-0.55),0.50,primary,25,True,SKY_BLUE)
    _textbox(slide,secondary_x,0.66,10.05-secondary_x,0.50,secondary,25,True,WHITE)
    _textbox(slide,0.47,1.23,9.25,0.25,strap,8.8,False,WHITE)

    # Left KPI panel
    panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.22), Inches(1.68), Inches(3.55), Inches(4.02))
    panel.fill.solid(); panel.fill.fore_color.rgb=NAVY_PANEL; panel.line.color.rgb=BLUE_LINE; panel.line.width=Pt(1.1)
    _textbox(slide,0.46,1.84,3.05,0.30,label.upper(),10.5,True,WHITE)
    card=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.35), Inches(2.18), Inches(3.28), Inches(1.52))
    card.fill.solid(); card.fill.fore_color.rgb=RGBColor(5,36,58); card.line.color.rgb=SKY_BLUE; card.line.width=Pt(1.2)
    circle=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), Inches(2.43), Inches(0.86), Inches(0.86))
    circle.fill.solid(); circle.fill.fore_color.rgb=RGBColor(8,45,67); circle.line.color.rgb=SKY_BLUE; circle.line.width=Pt(1.3)
    _textbox(slide,0.67,2.67,0.62,0.24,"KPI",10,True,SKY_BLUE,PP_ALIGN.CENTER)
    latest_fmt=("—" if latest is None else (_whole(latest) if unit=="percent" else _doi_whole(latest)))
    suffix="%" if unit=="percent" else " d"
    _textbox(slide,1.58,2.37,1.80,0.62,f"{latest_fmt}{suffix}",35,True,WHITE,PP_ALIGN.CENTER)
    if delta is not None:
        favorable=_delta_favorable(delta,good)
        col=AQUA if favorable is True else (RED if favorable is False else MUTED)
        arrow="▼" if float(delta)<0 else ("▲" if float(delta)>0 else "•")
        meaning="Positive" if favorable is True else ("Negative" if favorable is False else "Neutral")
        delta_fmt=_whole(abs(delta))
        _textbox(slide,1.32,3.27,2.15,0.26,f"{arrow} {delta_fmt} vs previous • {meaning}",8.6,True,col,PP_ALIGN.CENTER)

    # Summary panel
    summ=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.22), Inches(3.87), Inches(3.55), Inches(1.83))
    summ.fill.solid(); summ.fill.fore_color.rgb=NAVY_PANEL; summ.line.color.rgb=BLUE_LINE; summ.line.width=Pt(1.0)
    _textbox(slide,0.46,4.05,2.9,0.27,f"{period_title} SUMMARY",10.5,True,WHITE)
    numeric=[v for v in values if isinstance(v,(int,float))]
    if numeric:
        if unit=="percent":
            avg=_whole(sum(numeric)/len(numeric)); hi=_whole(max(numeric)); lo=_whole(min(numeric))
        else:
            avg=_doi_whole(sum(numeric)/len(numeric)); hi=_doi_whole(max(numeric)); lo=_doi_whole(min(numeric))
    else:
        avg=hi=lo="—"
    for i,(lab,val) in enumerate([("AVERAGE",avg),("HIGH",hi),("LOW",lo)]):
        x=0.40+i*1.06
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(4.43), Inches(0.94), Inches(0.92))
        sh.fill.solid(); sh.fill.fore_color.rgb=RGBColor(10,49,70); sh.line.color.rgb=BLUE_LINE
        _textbox(slide,x+0.07,4.60,0.80,0.18,lab,7.5,True,SKY_BLUE,PP_ALIGN.CENTER)
        _textbox(slide,x+0.03,4.91,0.88,0.30,f"{val}{suffix}",17,True,WHITE,PP_ALIGN.CENTER)

    # Right chart panel
    right=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.89), Inches(1.68), Inches(9.25), Inches(4.02))
    right.fill.solid(); right.fill.fore_color.rgb=NAVY_PANEL; right.line.color.rgb=BLUE_LINE; right.line.width=Pt(1.1)
    _textbox(slide,4.15,1.86,4.3,0.30,f"{period_title} • ACTUAL + TREND",12.5,True,SKY_BLUE)
    view="Month view • fits screen" if period_key=="ytd" else "Week view • fits screen"
    _textbox(slide,10.60,1.87,2.22,0.24,view,8,False,WHITE,PP_ALIGN.RIGHT)
    chart_stream=_kpi_chart_png(labels,values,trend,percent=(unit=="percent"),period=period_key,good=good)
    slide.shapes.add_picture(chart_stream, Inches(4.18), Inches(2.26), width=Inches(8.62), height=Inches(2.85))
    latest_dir="→"
    latest_meaning="Neutral"
    if delta is not None and abs(float(delta))>1e-12:
        latest_dir="↑" if float(delta)>0 else "↓"
        fav=_delta_favorable(delta,good)
        latest_meaning="Positive" if fav is True else ("Negative" if fav is False else "Neutral")
    _textbox(slide,4.17,5.24,8.60,0.22,f"Latest movement: {latest_dir} {latest_meaning}  •  Performance rule: {'higher DoI is better' if good=='high' else 'lower stock-out is better' if good=='low' else 'manage within healthy coverage'}",7.7,False,MUTED)

    # Bottom management action strip
    actions=_kpi_actions(label)
    for i,(num,head,body) in enumerate(actions):
        x=0.40+i*3.18
        if i>0:
            ln=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x-0.13), Inches(6.05), Inches(0.012), Inches(0.70))
            ln.fill.solid(); ln.fill.fore_color.rgb=RGBColor(88,139,173); ln.line.fill.background()
        circ=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(6.02), Inches(0.58), Inches(0.58))
        circ.fill.solid(); circ.fill.fore_color.rgb=NAVY; circ.line.color.rgb=SKY_BLUE; circ.line.width=Pt(1.4)
        _textbox(slide,x+0.08,6.19,0.42,0.18,num,8.5,True,SKY_BLUE,PP_ALIGN.CENTER)
        _textbox(slide,x+0.72,6.03,2.15,0.22,head,8.6,True,SKY_BLUE)
        _textbox(slide,x+0.72,6.34,2.14,0.40,body,8.2,False,WHITE)
    _textbox(slide,0.40,7.13,7.2,0.17,"RIGHT PRODUCTS.   RIGHT CUSTOMERS.   A STRONGER TOMORROW.",6.7,False,WHITE)
    _textbox(slide,9.45,7.13,3.45,0.17,"PEOPLE  •  PROCESS  •  PERFORMANCE",6.7,False,WHITE,PP_ALIGN.RIGHT)


def _textbox(slide, x, y, w, h, text, size=18, bold=False, color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.clear()
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = str(text); run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = color
    run.font.name = "Aptos Display" if size >= 20 else "Aptos"
    return box


def _add_logo(slide, x=9.70, y=0.32, w=2.85):
    if LOGO_PATH.exists():
        try:
            slide.shapes.add_picture(str(LOGO_PATH), Inches(x), Inches(y), width=Inches(w))
        except Exception:
            pass


def _title(slide, title, subtitle=None):
    # Executive blue header shared by all analytical slides.
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.12))
    band.fill.solid(); band.fill.fore_color.rgb = RGBColor(5, 22, 43); band.line.fill.background()
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.10), Inches(1.12))
    accent.fill.solid(); accent.fill.fore_color.rgb = BLUE; accent.line.fill.background()
    _add_logo(slide, x=10.25, y=0.24, w=2.38)
    _textbox(slide, 0.55, 0.28, 9.25, 0.44, title, 22, True, WHITE)
    if subtitle:
        _textbox(slide, 0.56, 0.76, 9.35, 0.24, subtitle, 8.6, False, MUTED)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.56), Inches(1.075), Inches(1.70), Inches(0.045))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    bar2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.28), Inches(1.075), Inches(0.82), Inches(0.045))
    bar2.fill.solid(); bar2.fill.fore_color.rgb = CYAN; bar2.line.fill.background()



def _section_divider(prs, blank, title: str, subtitle: str, items: List[str] | None = None):
    slide = prs.slides.add_slide(blank); _set_bg(slide, NAVY)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    band.fill.solid(); band.fill.fore_color.rgb = NAVY; band.line.fill.background()
    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.6), Inches(-1.45), Inches(6.2), Inches(6.2))
    glow.fill.solid(); glow.fill.fore_color.rgb = RGBColor(7, 52, 96); glow.line.fill.background()
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(1.18), Inches(0.10), Inches(4.85))
    rail.fill.solid(); rail.fill.fore_color.rgb = BLUE; rail.line.fill.background()
    _add_logo(slide, x=10.12, y=0.42, w=2.55)
    _textbox(slide, 0.98, 1.32, 8.9, 0.28, "MANAGEMENT REVIEW SECTION", 9, True, BLUE_PALE)
    _textbox(slide, 0.98, 1.85, 9.4, 0.72, title.upper(), 30, True, WHITE)
    _textbox(slide, 1.00, 2.70, 8.25, 0.42, subtitle, 14, False, MUTED)
    for i, item in enumerate(items or []):
        y = 3.62 + i * 0.54
        c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.00), Inches(y), Inches(0.30), Inches(0.30))
        c.fill.solid(); c.fill.fore_color.rgb = SKY_BLUE; c.line.fill.background()
        _textbox(slide, 1.48, y+0.02, 8.8, 0.18, item, 10.3, False, WHITE)
    _textbox(slide, 0.98, 6.86, 8.3, 0.18, "Designed for executive discussion, action ownership and branch-level follow-through.", 7.6, True, BLUE_PALE)
    return slide


def _mini_metric(slide, x, y, w, h, label, value, accent=BLUE_LINE, sub=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = DARK2; sh.line.color.rgb = RGBColor(30, 74, 118); sh.line.width = Pt(0.8)
    ac = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.055), Inches(h))
    ac.fill.solid(); ac.fill.fore_color.rgb = accent; ac.line.fill.background()
    _textbox(slide, x+0.15, y+0.10, w-0.25, 0.16, label, 7.1, True, MUTED)
    _textbox(slide, x+0.15, y+0.33, w-0.25, 0.22, value, 13.0, True, WHITE)
    if sub:
        _textbox(slide, x+0.15, y+h-0.20, w-0.25, 0.13, sub, 6.5, False, BLUE_PALE)
    return sh

def _card(slide, x, y, w, h, label, value, suffix="", delta=None, good="low"):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = DARK2; sh.line.color.rgb = RGBColor(30, 74, 118)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x+0.10), Inches(y+0.10), Inches(0.055), Inches(max(0.25, h-0.20)))
    accent.fill.solid(); accent.fill.fore_color.rgb = BLUE_LINE; accent.line.fill.background()
    _textbox(slide, x+0.22, y+0.12, w-0.42, 0.25, label, 8.5, True, MUTED)
    _textbox(slide, x+0.22, y+0.45, w-0.42, 0.42, f"{value}{suffix}", 22, True, WHITE)
    if delta is not None:
        favorable = _delta_favorable(delta, good)
        col = GREEN if favorable is True else (RED if favorable is False else MUTED)
        arrow = "▼" if delta < 0 else ("▲" if delta > 0 else "•")
        meaning = "" if favorable is None else (" • Positive" if favorable else " • Negative")
        _textbox(slide, x+0.18, y+h-0.38, w-0.36, 0.22, f"{arrow} {_whole(abs(delta))} vs previous{meaning}", 8, True, col)





def _table_header(slide, x, y, columns, widths, h=0.30):
    x0 = x
    for title, w in zip(columns, widths):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(10, 49, 91); sh.line.color.rgb = RGBColor(42, 91, 142)
        _textbox(slide, x0+0.04, y+0.07, w-0.08, h-0.08, title, 6.8, True, WHITE, PP_ALIGN.CENTER)
        x0 += w


def _table_row(slide, x, y, values, widths, h=0.28, danger=False):
    x0 = x
    for i, (val, w) in enumerate(zip(values, widths)):
        sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x0), Inches(y), Inches(w), Inches(h))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(8, 30, 56) if not danger else RGBColor(58, 28, 36)
        sh.line.color.rgb = RGBColor(30, 74, 118)
        col = RED if danger and i in (3,4,5,6) else WHITE
        align = PP_ALIGN.CENTER if i not in (1,2) else PP_ALIGN.LEFT
        _textbox(slide, x0+0.04, y+0.06, w-0.08, h-0.08, val, 6.5, i in (0,3,4,5,6), col, align)
        x0 += w


def _safe_pct(value):
    v = _whole(value)
    return "—" if v is None or v == "—" else f"{v}%"


def _branch_rows(all_branches: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    rows = []
    for b in all_branches or []:
        s = b.get("summary", {})
        rates = s.get("class_rates", {})
        rows.append({
            "branch": b.get("branch", ""),
            "area": b.get("area", ""),
            "class_a": rates.get("A", 0),
            "class_b": rates.get("B", 0),
            "class_c": rates.get("C", 0),
            "overall": s.get("overall_avg", 0),
        })
    rows.sort(key=lambda r: (-float(r.get("overall") or 0), -float(r.get("class_a") or 0), r.get("area", ""), r.get("branch", "")))
    return rows


def _add_kpi_summary_slide(prs, blank, data: Dict[str, Any]):
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "EXECUTIVE KPI SUMMARY", "Latest YTD and Weekly results across the six core Supply Chain KPIs.")
    cols = ["KPI", "YTD", "YTD Δ", "Weekly", "Weekly Δ", "Direction"]
    widths = [4.45, 1.25, 1.25, 1.25, 1.25, 1.50]
    _table_header(slide, 0.72, 1.35, cols, widths, 0.33)
    y = 1.76
    for name, kpi in data.get("kpis", {}).items():
        meta = kpi.get("meta", {})
        unit = "%" if meta.get("unit") == "percent" else " d"
        ytd = kpi.get("ytd", {})
        weekly = kpi.get("weekly", {})
        ytd_latest = ytd.get("latest")
        weekly_latest = weekly.get("latest")
        ytd_delta = ytd.get("delta")
        weekly_delta = weekly.get("delta")
        fmt_value = _whole if meta.get("unit") == "percent" else _doi_whole
        arrow = "→" if ytd_delta is None or abs(float(ytd_delta)) < 1e-12 else ("↑" if ytd_delta > 0 else "↓")
        favorable = _delta_favorable(ytd_delta, meta.get("good", "low"))
        direction = f"{arrow} " + ("Positive" if favorable is True else "Negative" if favorable is False else "Neutral")
        vals = [
            meta.get("label", name),
            "—" if ytd_latest is None else f"{fmt_value(ytd_latest)}{unit}",
            "—" if ytd_delta is None else f"{_whole(ytd_delta)}{unit}",
            "—" if weekly_latest is None else f"{fmt_value(weekly_latest)}{unit}",
            "—" if weekly_delta is None else f"{_whole(weekly_delta)}{unit}",
            direction,
        ]
        _table_row(slide, 0.72, y, vals, widths, 0.36, danger=(meta.get("unit") == "percent" and (ytd_latest or 0) > 20))
        y += 0.40
    _textbox(slide, 0.72, 6.62, 11.9, 0.26, "Use this page as the quick management checkpoint before reviewing YTD, Weekly, Area and all-Branch drilldowns.", 9.5, False, MUTED)


def _add_area_slide(prs, blank, area_data: Dict[str, Any]):
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "STOCK-OUT RATE SUMMARY · PER AREA", "Every Area shows Class A, Class B, Class C and Average Stock-Out Rate for executive action prioritization.")
    s = area_data.get("summary", {})
    cards = [("Class A Stock-Out", _whole(s.get("class_rates", {}).get("A", 0)), "%"), ("Class B Stock-Out", _whole(s.get("class_rates", {}).get("B", 0)), "%"), ("Class C Stock-Out", _whole(s.get("class_rates", {}).get("C", 0)), "%"), ("Overall Average", _whole(s.get("overall_avg", 0)), "%")]
    for i,(lab,val,suf) in enumerate(cards):
        _card(slide, 0.72+i*3.08, 1.42, 2.72, 1.08, lab, val, suf)
    ranking = area_data.get("ranking", [])
    _textbox(slide, 0.72, 2.86, 11.8, 0.25, "AREA CLASS A / B / C + AVERAGE STOCK-OUT RATE", 11, True, SKY_BLUE)
    cols=["#","Area","A","B","C","Overall"]
    widths=[0.45,2.3,0.9,0.9,0.9,1.1]
    _table_header(slide, 0.72, 3.20, cols, widths, 0.28)
    y=3.52
    for i,r in enumerate(ranking[:10], start=1):
        vals=[i, r.get('area',''), _safe_pct(r.get('class_a')), _safe_pct(r.get('class_b')), _safe_pct(r.get('class_c')), _safe_pct(r.get('overall_avg'))]
        _table_row(slide, 0.72, y, vals, widths, 0.27, danger=i<=3)
        y += 0.29
    _textbox(slide, 7.55, 2.86, 4.75, 0.25, "TOP BRANCH CLASS A STOCK-OUT", 11, True, SKY_BLUE)
    cols2=["#","Branch","Area","A %"]
    widths2=[0.45,2.45,1.05,0.80]
    _table_header(slide, 7.55, 3.20, cols2, widths2, 0.28)
    y=3.52
    for i,r in enumerate(area_data.get("branch_class_a", [])[:10], start=1):
        _table_row(slide, 7.55, y, [i, r.get('branch',''), r.get('area',''), _safe_pct(r.get('class_a'))], widths2, 0.27, danger=i<=5)
        y += 0.29


def _add_branch_class_a_ranking_slide(prs, blank, area_data: Dict[str, Any]):
    rows = area_data.get("branch_class_a", [])
    per_slide = 22
    for page, start in enumerate(range(0, len(rows), per_slide), start=1):
        slide = prs.slides.add_slide(blank); _set_bg(slide)
        total_pages = max(1, math.ceil(len(rows)/per_slide))
        _title(slide, f"BRANCH CLASS A STOCK-OUT RANKING · {page}/{total_pages}", "All branches are captured. Highest Class A exposure appears first.")
        cols=["#","Branch","Area","Class A %","SO / Status Count"]
        widths=[0.45,4.15,1.20,1.20,1.65]
        _table_header(slide, 0.72, 1.35, cols, widths, 0.30)
        y=1.70
        for i,r in enumerate(rows[start:start+per_slide], start=start+1):
            counts = r.get("class_a_counts", {}) or {}
            count_txt = f"{_whole(counts.get('stockout', 0))} / {_whole(counts.get('status_count', 0))}"
            vals=[i, r.get('branch',''), r.get('area',''), _safe_pct(r.get('class_a')), count_txt]
            _table_row(slide, 0.72, y, vals, widths, 0.23, danger=i<=10)
            y += 0.245


def _add_branch_all_slides(prs, blank, all_branches: List[Dict[str, Any]]):
    rows = _branch_rows(all_branches)
    per_slide = 17
    for page, start in enumerate(range(0, len(rows), per_slide), start=1):
        slide = prs.slides.add_slide(blank); _set_bg(slide)
        total_pages = max(1, math.ceil(len(rows)/per_slide))
        _title(slide, f"BRANCH STOCK-OUT RATE SUMMARY · {page}/{total_pages}", "Each Branch displays Class A, Class B, Class C and Average Stock-Out Rate for branch-level accountability.")
        cols=["#","Branch","Area","Class A %","Class B %","Class C %","Avg SO %"]
        widths=[0.45,3.9,1.2,1.0,1.0,1.0,1.1]
        _table_header(slide, 0.72, 1.35, cols, widths, 0.30)
        y=1.72
        for i,r in enumerate(rows[start:start+per_slide], start=start+1):
            vals=[i, r.get('branch',''), r.get('area',''), _safe_pct(r.get('class_a')), _safe_pct(r.get('class_b')), _safe_pct(r.get('class_c')), _safe_pct(r.get('overall'))]
            _table_row(slide, 0.72, y, vals, widths, 0.27, danger=i<=10)
            y += 0.30


def _add_priority_model_slide(prs, blank, all_branches: List[Dict[str, Any]]):
    # Capture all-branch model intelligence in one prioritized management list.
    model_rows=[]
    for b in all_branches or []:
        for cls in ("A","B","C"):
            for m in (b.get("classes", {}) or {}).get(cls, [])[:3]:
                status = m.get("stock_status", "") or ""
                risk = 2 if "stock" in status.lower() else (1 if "re" in status.lower() or "critical" in status.lower() else 0)
                model_rows.append({
                    "branch": b.get("branch", ""), "area": b.get("area", ""), "class": cls,
                    "rank": m.get("rank", 999999), "model": m.get("model", ""), "brand": m.get("brand", ""),
                    "status": status, "inventory": m.get("inventory", 0), "doi": m.get("doi", 0), "risk": risk,
                })
    model_rows.sort(key=lambda r: (r["class"] != "A", -r["risk"], r["rank"], r["area"], r["branch"]))
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "ALL-BRANCH MODEL PRIORITY SNAPSHOT", "Top model-level exceptions from branch Model Intelligence, prioritizing Class A and stock-risk conditions.")
    cols=["#","Area","Branch","Cls","Model","Status","Inv","DoI"]
    widths=[0.42,0.92,2.45,0.55,3.15,1.42,0.65,0.65]
    _table_header(slide, 0.72, 1.35, cols, widths, 0.30)
    y=1.70
    for i,r in enumerate(model_rows[:20], start=1):
        vals=[i, r['area'], r['branch'], r['class'], r['model'], r['status'], _whole(r['inventory']), _doi_whole(r['doi'])]
        _table_row(slide, 0.72, y, vals, widths, 0.245, danger=(r['risk']>0 and r['class']=='A'))
        y += 0.265


def _fit_text(draw, text: str, font, max_width: int) -> str:
    text = str(text or "")
    if draw.textbbox((0, 0), text, font=font)[2] <= max_width:
        return text
    suffix = "…"
    lo, hi = 0, len(text)
    while lo < hi:
        mid = (lo + hi + 1) // 2
        trial = text[:mid] + suffix
        if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
            lo = mid
        else:
            hi = mid - 1
    return text[:lo] + suffix


def _branch_model_table_png(rows: List[List[Any]], danger_flags: List[bool]) -> BytesIO:
    """Render a branch detail table as one PNG to avoid thousands of PPT shapes."""
    headers = ["#", "Brand", "Model", "Stock Status", "Inv", "Sug. Trf", "DoI"]
    width_ratios = [0.45, 1.25, 4.85, 1.80, 0.78, 0.95, 0.78]
    img_w = 1600
    header_h, row_h = 44, 38
    img_h = header_h + row_h * max(1, len(rows))
    image = Image.new("RGB", (img_w, img_h), (3, 12, 28))
    draw = ImageDraw.Draw(image)
    font_path = font_manager.findfont("DejaVu Sans")
    bold_path = font_manager.findfont("DejaVu Sans", fontext="ttf")
    font = ImageFont.truetype(font_path, 17)
    font_bold = ImageFont.truetype(bold_path, 17)
    header_font = ImageFont.truetype(bold_path, 17)
    total = sum(width_ratios)
    col_widths = [round(img_w * (w / total)) for w in width_ratios]
    col_widths[-1] += img_w - sum(col_widths)
    x_positions = [0]
    for w in col_widths:
        x_positions.append(x_positions[-1] + w)

    border = (30, 74, 118)
    header_bg = (10, 49, 91)
    row_bg = (8, 30, 56)
    danger_bg = (58, 28, 36)
    text_color = (248, 250, 252)
    danger_color = (248, 113, 113)

    for c, title in enumerate(headers):
        x0, x1 = x_positions[c], x_positions[c + 1]
        draw.rectangle((x0, 0, x1 - 1, header_h - 1), fill=header_bg, outline=border, width=1)
        box = draw.textbbox((0, 0), title, font=header_font)
        tx = x0 + max(8, (x1 - x0 - (box[2] - box[0])) // 2)
        ty = max(4, (header_h - (box[3] - box[1])) // 2 - 1)
        draw.text((tx, ty), title, font=header_font, fill=text_color)

    for r_idx, values in enumerate(rows):
        y0 = header_h + r_idx * row_h
        y1 = y0 + row_h
        is_danger = bool(danger_flags[r_idx]) if r_idx < len(danger_flags) else False
        for c, value in enumerate(values):
            x0, x1 = x_positions[c], x_positions[c + 1]
            draw.rectangle((x0, y0, x1 - 1, y1 - 1), fill=danger_bg if is_danger else row_bg, outline=border, width=1)
            align_left = c in (1, 2, 3)
            active_font = font_bold if c in (0, 2, 3, 4, 5, 6) else font
            shown = _fit_text(draw, value, active_font, max(10, x1 - x0 - 18))
            box = draw.textbbox((0, 0), shown, font=active_font)
            tw, th = box[2] - box[0], box[3] - box[1]
            tx = x0 + 9 if align_left else x0 + max(6, (x1 - x0 - tw) // 2)
            ty = y0 + max(3, (row_h - th) // 2 - 1)
            fill = danger_color if is_danger and c in (2, 3, 4, 5, 6) else text_color
            draw.text((tx, ty), shown, font=active_font, fill=fill)

    out = BytesIO()
    image.save(out, format="PNG", compress_level=1)
    out.seek(0)
    return out


def _add_branch_model_detail_slides(prs, blank, all_branches: List[Dict[str, Any]]):
    """Per-Branch Class A/B/C model pages for report-ready branch action.

    Each branch/class page now carries two layers of management context:
    1) Branch stock-out scorecard — Class A, B, C and Average Stock-Out Rate.
    2) Model action table — Stock Status, Inventory, Suggested Transfer and DoI.
    """
    branches = sorted(all_branches or [], key=lambda b: (b.get("area", ""), b.get("branch", "")))
    if not branches:
        slide = prs.slides.add_slide(blank); _set_bg(slide)
        _title(slide, "BRANCH MODEL INTELLIGENCE", "No Class A/B/C model rows available from the current Raw Distribution data.")
        return

    per_slide = 18
    for b in branches:
        branch = b.get("branch", "") or "—"
        area = b.get("area", "") or "—"
        classes = b.get("classes", {}) or {}
        summary = b.get("summary", {}) or {}
        rates = summary.get("class_rates", {}) or {}
        for cls in ("A", "B", "C"):
            rows = list(classes.get(cls, []) or [])
            if not rows:
                continue
            chunks = [rows[i:i+per_slide] for i in range(0, len(rows), per_slide)]
            total_pages = len(chunks)
            for page, chunk in enumerate(chunks, start=1):
                slide = prs.slides.add_slide(blank); _set_bg(slide)
                page_suffix = f" · {page}/{total_pages}" if total_pages > 1 else ""
                _title(slide, f"{branch} · CLASS {cls} MODEL ACTION PAGE{page_suffix}", f"{area} • Branch Stock-Out Rate plus model Stock Status, Inventory, Suggested Transfer and DoI.")

                # Branch stock-out rate summary on every branch model page.
                branch_cards = [
                    ("Class A SO", _safe_pct(rates.get("A", 0)), RED if float(rates.get("A", 0) or 0) >= 20 else BLUE_LINE),
                    ("Class B SO", _safe_pct(rates.get("B", 0)), RGBColor(245,158,11) if float(rates.get("B", 0) or 0) >= 20 else SKY_BLUE),
                    ("Class C SO", _safe_pct(rates.get("C", 0)), RGBColor(249,115,22) if float(rates.get("C", 0) or 0) >= 20 else BLUE_PALE),
                    ("Average SO", _safe_pct(summary.get("overall_avg", 0)), RED if float(summary.get("overall_avg", 0) or 0) >= 20 else GREEN),
                ]
                for i, (lab, val, accent) in enumerate(branch_cards):
                    _mini_metric(slide, 0.52 + i*2.42, 1.24, 2.17, 0.68, lab, val, accent)

                # Current class model metrics.
                risk_count = sum(1 for m in rows if any(x in str(m.get('stock_status','')).lower() for x in ['stock', 're', 'critical']))
                total_inv = sum(float(m.get('inventory') or 0) for m in rows)
                avg_doi = (sum(float(m.get('doi') or 0) for m in rows) / len(rows)) if rows else 0
                _textbox(slide, 10.35, 1.23, 2.35, 0.18, f"CLASS {cls} ACTION SUMMARY", 7.6, True, SKY_BLUE, PP_ALIGN.RIGHT)
                _textbox(slide, 10.35, 1.52, 2.35, 0.18, f"Models {len(rows)}  •  Risk {risk_count}  •  Inv {_whole(total_inv)}  •  Avg DoI {_doi_whole(avg_doi)}d", 7.0, False, MUTED, PP_ALIGN.RIGHT)

                table_rows = []
                danger_flags = []
                for i, m in enumerate(chunk, start=1 + (page-1)*per_slide):
                    status = str(m.get("stock_status", "") or "")
                    danger = (cls == "A" and any(x in status.lower() for x in ["stock", "re", "critical"]))
                    table_rows.append([
                        i,
                        m.get("brand", ""),
                        m.get("model", ""),
                        status,
                        _whole(m.get("inventory", 0)),
                        _whole(m.get("suggested_transfer", 0)),
                        _doi_whole(m.get("doi", 0)),
                    ])
                    danger_flags.append(danger)
                table_png = _branch_model_table_png(table_rows, danger_flags)
                slide.shapes.add_picture(table_png, Inches(0.72), Inches(2.25), width=Inches(11.58), height=Inches(4.78))
                _textbox(slide, 0.72, 7.10, 11.6, 0.18, "Report action: assign owner for Class A stock-out / reorder / critical items first, then rebalance inventory by DoI and Suggested Transfer.", 7.4, True, BLUE_PALE)



def _add_reorder_model_position_slides(prs, blank, reorder_card: Dict[str, Any], brand_filter: str = "All Brands"):
    rows = list((reorder_card or {}).get("rows", []) or [])
    if brand_filter and brand_filter != "All Brands":
        rows = [r for r in rows if (r.get("brand") or "Unspecified") == brand_filter]
    class_order = {"A": 0, "B": 1, "C": 2}
    rows.sort(key=lambda r: (class_order.get(r.get("class"), 9), int(r.get("rank", 999999)), r.get("brand", ""), r.get("model", "")))
    scope = brand_filter if brand_filter and brand_filter != "All Brands" else "All Brands"
    counts = {c: sum(1 for r in rows if r.get("class") == c) for c in ("A", "B", "C")}
    avgs = {}
    for c in ("A", "B", "C"):
        vals = [_doi_whole(r.get("doi", 0)) for r in rows if r.get("class") == c]
        vals = [v for v in vals if isinstance(v, (int, float))]
        avgs[c] = _doi_whole(sum(vals)/len(vals)) if vals else 0

    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "CLASS A / B / C MODEL POSITION", f"Imported Reorder priority • Brand filter: {scope} • DoI values follow the approved round-up rule.")
    cards=[("Class A",counts["A"],f"Avg DoI {avgs['A']} d"),("Class B",counts["B"],f"Avg DoI {avgs['B']} d"),("Class C",counts["C"],f"Avg DoI {avgs['C']} d"),("ABC Models",len(rows),scope)]
    for i,(lab,val,sub) in enumerate(cards):
        x=0.72+i*3.08
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.55), Inches(2.72), Inches(1.15))
        sh.fill.solid(); sh.fill.fore_color.rgb=DARK2; sh.line.color.rgb=RGBColor(51,65,85)
        _textbox(slide,x+0.18,1.73,2.3,0.22,lab,8.5,True,MUTED)
        _textbox(slide,x+0.18,2.02,2.3,0.34,val,20,True,WHITE)
        _textbox(slide,x+0.18,2.42,2.3,0.18,sub,7.5,True,SKY_BLUE)
    _textbox(slide,0.72,3.10,11.8,0.30,"Decision rule: prioritize Class A first; within each class, review lower DoI positions and imported Class Rank for replenishment action.",10.5,False,MUTED)
    if not rows:
        _textbox(slide,0.72,3.75,11.8,0.35,"No Reorder models match the selected Brand filter.",15,True,RED)
        return

    per_slide=20
    total_pages=max(1,math.ceil(len(rows)/per_slide))
    for page,start in enumerate(range(0,len(rows),per_slide),start=1):
        slide=prs.slides.add_slide(blank); _set_bg(slide)
        _title(slide,f"ABC MODEL POSITION · {page}/{total_pages}",f"Brand: {scope} • Sorted Class A → B → C, then imported Class Rank • DoI rounded up.")
        cols=["#","Class","Rank","Brand","Model","DoI","Stock Status"]
        widths=[0.45,0.70,0.65,1.45,4.35,0.70,2.55]
        _table_header(slide,0.72,1.35,cols,widths,0.30)
        y=1.70
        for i,r in enumerate(rows[start:start+per_slide],start=start+1):
            danger=(r.get("class")=="A" and ("stock" in str(r.get("stock_status","")).lower() or _doi_whole(r.get("doi",0)) <= 3))
            vals=[i,r.get("class","—"),r.get("rank","—"),r.get("brand","—"),r.get("model","—"),f"{_doi_whole(r.get('doi',0))} d",r.get("stock_status","—")]
            _table_row(slide,0.72,y,vals,widths,0.245,danger=danger)
            y+=0.265



def _fmt_qty(value: Any) -> str:
    try:
        return f"{int(round(float(value or 0))):,}"
    except Exception:
        return "0"


def _fmt_money_short(value: Any) -> str:
    try:
        v = float(value or 0)
    except Exception:
        return "₱0"
    if abs(v) >= 1_000_000:
        return f"₱{v/1_000_000:.1f}M"
    if abs(v) >= 1_000:
        return f"₱{v/1_000:.0f}K"
    return f"₱{v:,.0f}"


def _add_aging_executive_slides(prs, blank, summary: Dict[str, Any] | None):
    """Add management-ready Aging slides when the integrated module has data."""
    if not summary:
        return

    # Slide 1 — executive snapshot.
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    as_of = str(summary.get("as_of_date") or "—")
    _title(slide, "MOTORCYCLE AGING · EXECUTIVE SNAPSHOT", f"Integrated Aging Decision Intelligence • As of {as_of}")

    cards = [
        ("TOTAL UNITS", _fmt_qty(summary.get("total_qty")), BLUE_LINE),
        ("AVG AGE", f"{_whole(summary.get('avg_age', 0))} d", SKY_BLUE),
        ("91+ DAYS", f"{_fmt_qty(summary.get('aged_90'))} • {_whole(summary.get('aged_90_pct'))}%", RGBColor(249, 115, 22)),
        ("366+ DAYS", f"{_fmt_qty(summary.get('aged_365'))} • {_whole(summary.get('aged_365_pct'))}%", RED),
        ("91+ VALUE", _fmt_money_short(summary.get("aged_90_value")), RGBColor(245, 158, 11)),
        ("OLDEST UNIT", f"{_whole(summary.get('oldest', 0))} d", RED),
    ]
    x0=0.62
    for i,(label,value,accent_col) in enumerate(cards):
        x=x0+(i%3)*2.90
        y=1.42+(i//3)*1.08
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.68), Inches(0.88))
        sh.fill.solid(); sh.fill.fore_color.rgb=NAVY_PANEL; sh.line.color.rgb=RGBColor(30,74,118)
        ac=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(0.88))
        ac.fill.solid(); ac.fill.fore_color.rgb=accent_col; ac.line.fill.background()
        _textbox(slide,x+0.20,y+0.13,2.25,0.18,label,7.8,True,MUTED)
        _textbox(slide,x+0.20,y+0.40,2.30,0.30,value,16.5,True,WHITE)

    # Age distribution — deliberately semantic, not all-blue.
    _textbox(slide,0.66,3.82,6.1,0.28,"AGE DISTRIBUTION",12,True,WHITE)
    labels=list(summary.get("bucket_labels") or [])
    vals=[float(v or 0) for v in (summary.get("bucket_values") or [])]
    bucket_colors=[BLUE_LINE, SKY_BLUE, RGBColor(245,158,11), RED]
    maxv=max(vals) if vals else 1
    for i,(lab,val) in enumerate(zip(labels, vals)):
        y=4.25+i*0.52
        _textbox(slide,0.66,y,1.23,0.20,lab,7.8,True,MUTED)
        bg=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.93), Inches(y+0.01), Inches(3.70), Inches(0.22))
        bg.fill.solid(); bg.fill.fore_color.rgb=RGBColor(18,42,70); bg.line.fill.background()
        w=max(0.06, 3.70*(val/maxv if maxv else 0)) if val else 0.06
        bar=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.93), Inches(y+0.01), Inches(w), Inches(0.22))
        bar.fill.solid(); bar.fill.fore_color.rgb=bucket_colors[min(i,len(bucket_colors)-1)]; bar.line.fill.background()
        _textbox(slide,5.79,y-0.01,0.80,0.24,_fmt_qty(val),8.5,True,WHITE,PP_ALIGN.RIGHT)

    # Risk interpretation.
    risk=str(summary.get("risk_level") or "Controlled")
    risk_color=RED if risk.lower()=="high" else (RGBColor(245,158,11) if risk.lower()=="watch" else GREEN)
    panel=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.05), Inches(3.80), Inches(5.55), Inches(2.45))
    panel.fill.solid(); panel.fill.fore_color.rgb=RGBColor(8,30,56); panel.line.color.rgb=RGBColor(30,74,118)
    _textbox(slide,7.30,4.07,2.5,0.22,"AGING RISK SIGNAL",8.5,True,MUTED)
    pill=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(10.58), Inches(3.98), Inches(1.65), Inches(0.40))
    pill.fill.solid(); pill.fill.fore_color.rgb=risk_color; pill.line.fill.background()
    _textbox(slide,10.66,4.10,1.48,0.16,risk.upper(),8.5,True,NAVY,PP_ALIGN.CENTER)
    _textbox(slide,7.30,4.52,4.85,0.90,summary.get("risk_message") or "",11,False,WHITE)
    _textbox(slide,7.30,5.55,4.85,0.22,f"Healthy inventory: {_whole(summary.get('healthy_pct',0))}% • 91+ aged value: {_whole(summary.get('aged_value_pct',0))}%",8.5,True,BLUE_PALE)
    _textbox(slide,7.30,5.89,4.85,0.20,f"Coverage: {_fmt_qty(summary.get('branch_count'))} branches • {_fmt_qty(summary.get('area_count'))} areas",8,False,MUTED)

    # Slide 2 — action priorities.
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "MOTORCYCLE AGING · ACTION PRIORITIES", "Ranked 91+ day exposure by Area, Branch and Model")
    tops=[
        ("TOP RISK AREA", summary.get("top_area") or {}, "name"),
        ("TOP RISK BRANCH", summary.get("top_branch") or {}, "name"),
        ("TOP AGED MODEL", summary.get("top_model") or {}, "standard_description"),
    ]
    for i,(label,obj,key) in enumerate(tops):
        x=0.65+i*4.18
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.42), Inches(3.92), Inches(1.20))
        sh.fill.solid(); sh.fill.fore_color.rgb=NAVY_PANEL; sh.line.color.rgb=RGBColor(30,74,118)
        _textbox(slide,x+0.20,1.60,3.50,0.18,label,8,True,MUTED)
        _textbox(slide,x+0.20,1.88,3.50,0.26,obj.get(key,"—") or "—",13,True,WHITE)
        _textbox(slide,x+0.20,2.27,3.50,0.18,f"91+ {_fmt_qty(obj.get('aged_90'))} • {_whole(obj.get('aged_90_pct',0))}% • Oldest {_whole(obj.get('oldest',0))}d",7.3,True,RGBColor(248,180,120))

    # Area table.
    _textbox(slide,0.68,2.94,5.95,0.25,"TOP AREAS BY 91+ DAY UNITS",10.5,True,WHITE)
    aw=[2.15,0.90,0.90,1.00,0.80]
    _table_header(slide,0.68,3.30,["Area","Units","91+","91+ %","Oldest"],aw,0.32)
    y=3.66
    for r in (summary.get("area_ranking") or [])[:6]:
        _table_row(slide,0.68,y,[r.get("name","—"),_fmt_qty(r.get("qty")),_fmt_qty(r.get("aged_90")),f"{_whole(r.get('aged_90_pct',0))}%",f"{_whole(r.get('oldest',0))}d"],aw,0.32,danger=float(r.get("aged_90_pct") or 0)>=40)
        y+=0.36

    # Branch table.
    _textbox(slide,7.08,2.94,5.55,0.25,"TOP BRANCHES BY 91+ DAY UNITS",10.5,True,WHITE)
    bw=[2.15,1.45,0.75,0.85,0.75]
    _table_header(slide,7.08,3.30,["Branch","Area","91+","91+ %","Oldest"],bw,0.32)
    y=3.66
    for r in (summary.get("branch_ranking") or [])[:6]:
        _table_row(slide,7.08,y,[r.get("name","—"),r.get("area","—"),_fmt_qty(r.get("aged_90")),f"{_whole(r.get('aged_90_pct',0))}%",f"{_whole(r.get('oldest',0))}d"],bw,0.32,danger=float(r.get("aged_90_pct") or 0)>=40)
        y+=0.36

    _textbox(slide,0.68,6.62,11.9,0.30,"Management action: prioritize transfer/sell-through for concentrated 91+ exposure, then review 181+ and 366+ units for liquidation or recovery decisions.",9.2,True,BLUE_PALE)



def _add_report_agenda_slide(prs, blank, has_aging: bool):
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _title(slide, "REPORT ROADMAP", "Customized management-ready flow for KPI review, stock-out exposure, branch model actions and Aging Unit risk.")
    steps = [
        ("01", "KPI PERFORMANCE", "YTD and Weekly trend pages with blue Actual and Trend lines."),
        ("02", "STOCK-OUT RATE SUMMARY", "Per Area and per Branch Class A/B/C and Average Stock-Out Rate."),
        ("03", "BRANCH MODEL ACTIONS", "Every branch shows Class A/B/C model lists with Stock Status, Inventory and DoI."),
        ("04", "ORDER / REORDER INTELLIGENCE", "ABC model position and management replenishment review."),
    ]
    if has_aging:
        steps.append(("05", "AGING UNIT DECISION INTELLIGENCE", "Aging exposure, 91+ risk and unit-level priorities."))
    for i,(num,head,body) in enumerate(steps):
        y=1.45+i*0.95
        c=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.78), Inches(y), Inches(0.62), Inches(0.62))
        c.fill.solid(); c.fill.fore_color.rgb=BLUE; c.line.color.rgb=SKY_BLUE
        _textbox(slide,0.88,y+0.20,0.42,0.16,num,8,True,WHITE,PP_ALIGN.CENTER)
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.65), Inches(y-0.05), Inches(10.8), Inches(0.72))
        sh.fill.solid(); sh.fill.fore_color.rgb=DARK2; sh.line.color.rgb=RGBColor(30,74,118)
        _textbox(slide,1.90,y+0.07,3.35,0.22,head,10.5,True,SKY_BLUE)
        _textbox(slide,5.15,y+0.08,6.90,0.22,body,9.2,False,WHITE)
    _textbox(slide,0.78,6.75,11.3,0.24,"Use this deck directly in the report: lead with KPI movement, quantify Area/Branch stock-out exposure, then assign action by model and Aging risk.",9.2,True,BLUE_PALE)

def build_presentation(data: Dict[str, Any], area_data: Dict[str, Any], branch_data: Dict[str, Any], all_branches: List[Dict[str, Any]] | None = None, reorder_brand: str = "All Brands") -> bytes:
    """Build the management deck.

    v2.45: Export is a highly customized report-ready executive deck with KPI trends, Area/Branch Class A-B-C Stock-Out summaries, Branch model action pages and Aging Unit decision intelligence.
    """
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1 Cover — native vector executive blue design.
    slide = prs.slides.add_slide(blank); _set_bg(slide, NAVY)
    rail = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.35), Inches(0), Inches(2.983), Inches(7.5))
    rail.fill.solid(); rail.fill.fore_color.rgb = RGBColor(6, 35, 68); rail.line.fill.background()
    rail2 = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.35), Inches(0), Inches(0.12), Inches(7.5))
    rail2.fill.solid(); rail2.fill.fore_color.rgb = BLUE; rail2.line.fill.background()
    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(10.35), Inches(0.10))
    top.fill.solid(); top.fill.fore_color.rgb = CYAN; top.line.fill.background()
    _add_logo(slide, x=10.68, y=0.42, w=2.05)
    _textbox(slide, 0.72, 0.72, 5.6, 0.28, "SCM • INVENTORY CONTROL TOWER", 10, True, BLUE_PALE)
    _textbox(slide, 0.72, 1.48, 8.9, 0.66, "INVENTORY &", 34, True, WHITE)
    _textbox(slide, 0.72, 2.15, 8.9, 0.66, "DISTRIBUTION PLANNING", 34, True, SKY_BLUE)
    _textbox(slide, 0.72, 2.82, 8.9, 0.48, "EXECUTIVE PERFORMANCE REVIEW", 20, True, WHITE)
    _textbox(slide, 0.74, 3.48, 7.8, 0.30, "Visibility • Availability • Replenishment • Distribution", 11, False, MUTED)
    # Date capsule always reflects the actual generation date.
    capsule = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(4.38), Inches(3.45), Inches(0.70))
    capsule.fill.solid(); capsule.fill.fore_color.rgb = RGBColor(8, 37, 71); capsule.line.color.rgb = BLUE_LINE; capsule.line.width = Pt(1.0)
    _textbox(slide, 0.94, 4.52, 3.0, 0.18, "GENERATED", 8, True, MUTED)
    _textbox(slide, 0.94, 4.75, 3.0, 0.23, _generated_date(data), 13.5, True, WHITE)
    # Management pillars.
    pillars=[("01","VISIBILITY","One source of truth"),("02","CONTROL","Actionable exceptions"),("03","SERVICE","Branch availability")]
    for i,(num,head,sub) in enumerate(pillars):
        y=4.48+i*0.86
        c=slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.77), Inches(y), Inches(0.50), Inches(0.50))
        c.fill.solid(); c.fill.fore_color.rgb=BLUE; c.line.fill.background()
        _textbox(slide,10.87,y+0.14,0.30,0.14,num,7.5,True,WHITE,PP_ALIGN.CENTER)
        _textbox(slide,11.45,y+0.01,1.35,0.20,head,8.5,True,WHITE)
        _textbox(slide,11.45,y+0.25,1.35,0.22,sub,7.0,False,MUTED)
    _textbox(slide, 0.72, 6.92, 8.7, 0.22, "SUPPLY CHAIN MANAGEMENT • PEOPLE • PROCESS • PERFORMANCE", 7.4, True, BLUE_PALE)

    _add_report_agenda_slide(prs, blank, bool(data.get("aging_summary")))

    _section_divider(prs, blank, "KPI Performance", "Trend movement and management interpretation for YTD and Weekly KPI results.", ["Executive KPI summary", "YTD KPI trend pages", "Weekly KPI trend pages"])
    _add_kpi_summary_slide(prs, blank, data)

    # Approved KPI template: one full executive slide per KPI for both YTD and Weekly views.
    for period_key, period_title in [("ytd", "YTD"), ("weekly", "WEEKLY")]:
        for name, kpi in data.get("kpis", {}).items():
            _add_kpi_template_slide(prs, blank, name, kpi, period_key, period_title)

    _section_divider(prs, blank, "Stock-Out Rate Intelligence", "Per Area and per Branch Class A, B, C and Average Stock-Out Rate for decision accountability.", ["Area Class A / B / C + Average Stock-Out Rate", "Branch Class A / B / C + Average Stock-Out Rate", "Class A exposure ranking"])
    _add_area_slide(prs, blank, area_data)
    _add_branch_all_slides(prs, blank, all_branches or ([branch_data] if branch_data else []))
    _add_branch_class_a_ranking_slide(prs, blank, area_data)
    _section_divider(prs, blank, "Branch Model Action Pages", "Every Branch/Class page includes Stock Status, Inventory, Suggested Transfer and DoI.", ["Class A, Class B and Class C separated per branch", "Stock Status + Inventory + DoI visible for every model", "Branch Stock-Out Rate scorecard appears on each model page"])
    _add_branch_model_detail_slides(prs, blank, all_branches or ([branch_data] if branch_data else []))
    _add_priority_model_slide(prs, blank, all_branches or ([branch_data] if branch_data else []))
    _section_divider(prs, blank, "Management Order Intelligence", "ABC model position and reorder review for replenishment planning after branch risk review.", ["Class A / B / C model position", "Imported Class Rank and DoI", "Brand-filtered order planning support"])
    _add_reorder_model_position_slides(prs, blank, data.get("reorder_card", {}), reorder_brand)
    if data.get("aging_summary"):
        _section_divider(prs, blank, "Aging Unit Decision Intelligence", "Integrated Motorcycle Aging exposure, 91+ risk and action priorities from the same consolidated workbook.", ["Aging Unit executive snapshot", "91+ day unit and capital exposure", "Area, Branch and Model aging priorities"])
    _add_aging_executive_slides(prs, blank, data.get("aging_summary"))

    # Closing slide
    slide = prs.slides.add_slide(blank); _set_bg(slide)
    _textbox(slide, 0.72, 1.0, 11.9, 0.55, "FROM VISIBILITY TO ACTION", 28, True, WHITE)
    _textbox(slide, 0.72, 1.65, 11.6, 0.4, "Translate KPI movement, stock exposure and Motorcycle Aging into clear Area, Branch and Model actions using one management-ready review deck.", 15, False, MUTED)
    actions = [
        "1  Review YTD and Weekly KPI movement before approving corrective actions.",
        "2  Review Class A/B/C Model Position, starting with Class A Rank and low DoI exceptions.",
        "3  Prioritize high stock-out Areas and Class A branch exposure.",
        "4  Use all-Branch rankings to assign follow-up owners and replenishment priority.",
        "5  Review Motorcycle Aging priorities, then re-import current source files and regenerate the deck for each management review.",
    ]
    for i,a in enumerate(actions):
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(2.45+i*0.72), Inches(11.8), Inches(0.55))
        sh.fill.solid(); sh.fill.fore_color.rgb=DARK2; sh.line.color.rgb=RGBColor(51,65,85)
        _textbox(slide, 1.0, 2.60+i*0.72, 11.2, 0.24, a, 13, i==0, SKY_BLUE if i==0 else WHITE)

    out = BytesIO(); prs.save(out); return out.getvalue()
